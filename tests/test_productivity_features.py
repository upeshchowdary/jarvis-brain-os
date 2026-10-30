"""Phase 5 features: time parsing, scheduler, store, routines, file ops + undo, documents, health, triggers."""

import os
import zipfile
from datetime import datetime
from unittest.mock import AsyncMock, MagicMock

import pytest

from automation.gateway import ActionRequest, gateway
from core.config import settings
from services.timeparse import next_occurrence, parse_when, strip_time_phrase
from tools import ToolContext, load_builtin_tools

FRI_2PM = datetime(2026, 10, 30, 14, 0)
registry = load_builtin_tools()


@pytest.fixture(autouse=True)
def relaxed(monkeypatch):
    gateway.resume()
    yield
    gateway.resume()


# ── Time parsing ─────────────────────────────────────────────────────────────
@pytest.mark.parametrize("text,due,rec", [
    ("at 6 pm to submit the report", datetime(2026, 10, 30, 18, 0), None),
    ("in 10 minutes", datetime(2026, 10, 30, 14, 10), None),
    ("in an hour and a half", datetime(2026, 10, 30, 15, 30), None),
    ("tomorrow at 9 call mom", datetime(2026, 10, 31, 9, 0), None),
    ("every weekday at 9", datetime(2026, 11, 2, 9, 0), "weekdays"),
    ("every monday at 10:15", datetime(2026, 11, 2, 10, 15), "weekly:0"),
    ("tonight at 8", datetime(2026, 10, 30, 20, 0), None),
    ("every 30 minutes", datetime(2026, 10, 30, 14, 30), "every:1800"),
    ("at 18:30", datetime(2026, 10, 30, 18, 30), None),
    ("every day at 7am", datetime(2026, 10, 31, 7, 0), "daily"),
    ("next monday 10am", datetime(2026, 11, 2, 10, 0), None),
])
def test_parse_when(text, due, rec):
    p = parse_when(text, FRI_2PM)
    assert p is not None and p.due == due and p.recurrence == rec


def test_parse_when_none_and_label_cleanup():
    assert parse_when("the meeting", FRI_2PM) is None
    p = parse_when("6 pm to submit the report", FRI_2PM)
    assert strip_time_phrase("6 pm to submit the report", p) == "submit the report"


def test_next_occurrence():
    assert next_occurrence("daily", datetime(2026, 10, 30, 7), datetime(2026, 10, 30, 8)) == datetime(2026, 10, 31, 7)
    assert next_occurrence("weekdays", datetime(2026, 10, 30, 9), datetime(2026, 10, 30, 9, 1)) == datetime(2026, 11, 2, 9)
    assert next_occurrence("every:600", datetime(2026, 10, 30, 9), datetime(2026, 10, 30, 9, 25)) == datetime(2026, 10, 30, 9, 30)


# ── Scheduler & store ────────────────────────────────────────────────────────
async def test_scheduler_fires_and_reschedules(monkeypatch):
    import services.scheduler as sched
    from services.store import store

    notify = AsyncMock()
    monkeypatch.setattr(sched, "notify", notify)
    now = datetime.now().timestamp()
    one = await store.add_reminder("drink water", now - 5)
    rec = await store.add_reminder("stretch", now - 5, "every:600")
    fired = await sched.scheduler.tick(now)
    assert fired >= 2 and notify.await_count >= 2
    remaining = {r["id"]: r for r in await store.reminders()}
    assert one not in remaining and rec in remaining and remaining[rec]["due_ts"] > now
    await store.cancel_reminders("stretch")


async def test_reminder_tool_and_command_route():
    from brain.pipeline import brain

    res = await brain.handle("remind me at 6 pm to submit the report", "rem-test")
    assert res.route == "fast_path" and "submit the report" in res.response
    listed = await brain.handle("what are my reminders", "rem-test")
    assert "submit the report" in listed.response
    cancelled = await brain.handle("cancel reminder submit the report", "rem-test")
    assert "Cancelled 1" in cancelled.response


async def test_notes_and_todos_via_commands():
    from brain.pipeline import brain

    await brain.handle("add renew passport to my to-do list", "todo-test")
    assert "renew passport" in (await brain.handle("what's on my to-do list", "todo-test")).response
    assert "Done" in (await brain.handle("mark renew passport as done", "todo-test")).response
    await brain.handle("take a note the wifi password is on the router", "todo-test")
    assert "wifi password" in (await brain.handle("show my notes", "todo-test")).response


# ── Routines ─────────────────────────────────────────────────────────────────
def test_routine_plan_variables_and_conditions():
    from services import routines

    r = routines.find("study")
    acts, local = routines.plan(r, {"minutes": 25}, status={"battery_percent": 10, "processes": set()}, now=FRI_2PM)
    assert [a.op for a in acts] == ["app.open", "browser.open_url", "system.volume"]
    says = [s["say"] for s in local if "say" in s]
    assert any("25 minutes" in s for s in says) and any("battery is low" in s.lower() for s in says)
    _, local2 = routines.plan(r, {}, status={"battery_percent": 90, "processes": set()}, now=FRI_2PM)
    assert not any("battery" in s.get("say", "").lower() for s in local2)


async def test_routine_runs_as_single_confirmed_plan(monkeypatch):
    from services import routines

    monkeypatch.setattr(routines, "_status", AsyncMock(return_value={"battery_percent": 80, "processes": set()}))
    result, msgs = await routines.run("shutdown routine", "routine-s")
    assert result.status == "needs_confirmation" and "Shut down" in result.preview and "Close chrome" in result.preview
    gateway.cancel(result.pending_id)
    dry, msgs = await routines.run("study mode", "routine-d", dry_run=True)
    assert dry.status == "dry_run"


# ── File operations with undo ────────────────────────────────────────────────
async def test_organize_and_undo(tmp_path, monkeypatch):
    monkeypatch.setattr(settings, "CONFIRM_MEDIUM_RISK", False)
    for name in ("a.pdf", "b.jpg", "c.zip", "notes.txt", "d.py"):
        (tmp_path / name).write_text("x")
    res = await gateway.execute(ActionRequest("files.organize", {"folder": str(tmp_path)}), "org")
    assert res.status == "needs_confirmation" and "Documents" in res.preview
    done = await gateway.confirm(res.pending_id, "org")
    assert done.status == "executed"
    assert (tmp_path / "Documents" / "a.pdf").exists() and (tmp_path / "Images" / "b.jpg").exists()
    undo = await gateway.undo_last("org")
    assert undo.status == "executed" and (tmp_path / "a.pdf").exists() and (tmp_path / "b.jpg").exists()


async def test_create_rename_trash(tmp_path, monkeypatch):
    monkeypatch.setattr(settings, "CONFIRM_MEDIUM_RISK", False)
    f = tmp_path / "draft.txt"
    assert (await gateway.execute(ActionRequest("file.create", {"path": str(f), "content": "hi"}), "fs")).status == "executed"
    assert f.read_text() == "hi"
    assert (await gateway.execute(ActionRequest("file.rename", {"src": str(f), "new_name": "final.txt"}), "fs")).status == "executed"
    assert (tmp_path / "final.txt").exists()
    trashed = MagicMock()
    import send2trash

    monkeypatch.setattr(send2trash, "send2trash", trashed)
    res = await gateway.execute(ActionRequest("file.trash", {"paths": [str(tmp_path / "final.txt")]}), "fs")
    assert res.status == "needs_confirmation"  # deleting is HIGH even with medium relaxed
    await gateway.confirm(res.pending_id, "fs")
    trashed.assert_called_once()


async def test_protected_locations_refused():
    res = await gateway.execute(ActionRequest("file.create", {"path": "C:/Windows/evil.txt", "content": "x"}), "prot", confirmed=True)
    assert res.status == "failed" and "system location" in res.message
    res = await gateway.execute(ActionRequest("file.create", {"path": "/etc/evil.conf", "content": "x"}), "prot", confirmed=True)
    assert res.status == "failed"


@pytest.mark.parametrize("path", ["D:\\Windows\\evil.txt", "\\Windows\\System32\\evil.txt", "C:\\Program Files\\evil.txt",
                                  "c:/programdata/evil.txt", "/etc/evil.conf", "\\etc\\evil.conf"])
async def test_system_roots_refused_in_every_spelling(path, tmp_path, monkeypatch):
    """Rooted paths without a drive ("\\Windows\\…", "/etc/…") aren't "absolute" on Windows; they must still be
    treated as system locations rather than resolved like natural language."""
    monkeypatch.chdir(tmp_path)
    res = await gateway.execute(ActionRequest("file.create", {"path": path, "content": "x"}), "prot", confirmed=True)
    assert res.status == "failed" and "system location" in res.message, (path, res.message)
    assert not any(tmp_path.iterdir()), "nothing may be written anywhere"


def test_rooted_paths_resolve_to_absolute():
    from automation.paths import resolve

    for raw in ("/etc/evil.conf", "\\Windows\\evil.txt"):  # rooted: never handed to the natural-language resolver
        assert resolve(raw).is_absolute(), raw


async def test_unzip_blocks_zip_slip(tmp_path, monkeypatch):
    monkeypatch.setattr(settings, "CONFIRM_MEDIUM_RISK", False)
    evil = tmp_path / "evil.zip"
    with zipfile.ZipFile(evil, "w") as zf:
        zf.writestr("../../escape.txt", "pwned")
    res = await gateway.execute(ActionRequest("files.unzip", {"archive": str(evil), "dest": str(tmp_path / "out")}), "zip")
    assert res.status == "failed" and not (tmp_path.parent / "escape.txt").exists()


async def test_bulk_rename_preview_and_duplicates(tmp_path, monkeypatch):
    monkeypatch.setattr(settings, "CONFIRM_MEDIUM_RISK", False)
    for i in range(3):
        (tmp_path / f"IMG_{i}.jpg").write_bytes(b"same")
    res = await gateway.execute(ActionRequest("files.bulk_rename", {"folder": str(tmp_path), "template": "goa_{n}"}), "br")
    assert res.status == "executed" and (tmp_path / "goa_001.jpg").exists()
    dups = await registry.get("find_duplicate_files").run(ToolContext(), folder=str(tmp_path))
    assert "3×" in dups.message


# ── Documents ────────────────────────────────────────────────────────────────
async def test_documents_ask_bm25(tmp_path):
    (tmp_path / "physics.md").write_text("# Notes\nNewton's second law: force equals mass times acceleration.\n")
    (tmp_path / "history.txt").write_text("The French revolution began in 1789.")
    res = await registry.get("documents_ask").run(ToolContext(), folder=str(tmp_path), question="what is newton's second law about force?")
    assert res.success and "physics.md" in res.content_for_llm.split("\n")[0]


async def test_extract_tables_and_convert(tmp_path, monkeypatch):
    monkeypatch.setattr(settings, "CONFIRM_MEDIUM_RISK", False)
    md = tmp_path / "marks.md"
    md.write_text("| name | score |\n|---|---|\n| Asha | 91 |\n| Ravi | 85 |\n")
    res = await gateway.execute(ActionRequest("docs.extract_tables", {"path": str(md)}), "docs")
    assert res.status == "executed"
    from openpyxl import load_workbook

    ws = load_workbook(tmp_path / "marks_tables.xlsx").active
    assert [c.value for c in ws[2]] == ["Asha", "91"]
    res = await gateway.execute(ActionRequest("docs.convert", {"path": str(md), "to": "docx"}), "docs")
    assert res.status == "executed" and (tmp_path / "marks.docx").exists()


# ── Health monitor & triggers ────────────────────────────────────────────────
def test_health_evaluate():
    from services.health import evaluate

    alerts = evaluate({"battery_percent": 9, "battery_plugged": False, "cpu_percent": 99, "ram_percent": 50, "disk_percent": 97,
                       "disk_free_gb": 3}, [{"name": "chrome.exe"}])
    kinds = {k for k, _ in alerts}
    assert {"battery", "cpu", "disk"} <= kinds and "ram" not in kinds
    assert "kill chrome.exe" in dict(alerts)["cpu"]


def test_time_trigger():
    from services.triggers import time_trigger_due

    trig = {"type": "time", "at": "09:00", "days": ["mon", "tue", "wed", "thu", "fri"]}
    assert time_trigger_due(trig, datetime(2026, 11, 2, 9, 0), None)
    assert not time_trigger_due(trig, datetime(2026, 11, 2, 9, 0), "2026-11-02")
    assert not time_trigger_due(trig, datetime(2026, 11, 1, 9, 0), None)  # Sunday


async def test_file_created_trigger_replays_workflow(tmp_path, monkeypatch):
    from automation.workflow_recorder import Workflow, WorkflowStep, workflow_recorder
    from services import triggers

    wf = Workflow(name="sort pdfs", steps=[WorkflowStep(0, "wait", value="0")],
                  triggers=[{"type": "file_created", "folder": str(tmp_path), "pattern": "*.pdf"}])
    workflow_recorder.save(wf)
    fire = AsyncMock()
    tm = triggers.TriggerManager()
    monkeypatch.setattr(tm, "_fire", fire)
    monkeypatch.setattr(tm, "_snapshot", lambda: (set(), set()))
    assert await tm.check_once() == []  # first pass primes the folder snapshot
    (tmp_path / "invoice.pdf").write_text("x")
    assert await tm.check_once() == ["sort_pdfs"]
    fire.assert_awaited_once()
    workflow_recorder.delete_workflow("sort pdfs")


# ── Tool selection, research, dev ───────────────────────────────────────────
def test_tool_selection_keeps_prompt_small():
    from tools.selection import select_tools

    general = [t.name for t in select_tools("what is quantum computing", registry.read_only())]
    assert len(general) <= 12 and "web_search" in general and "set_reminder" not in general
    files = [t.name for t in select_tools("organize my downloads folder", registry.all())]
    assert "organize_folder" in files and len(files) <= 28


async def test_research_combines_sources(monkeypatch):
    import tools.builtin.research as research

    monkeypatch.setattr(research, "ddg_search", AsyncMock(return_value=[
        {"title": "A", "snippet": "sa", "url": "https://a.example"}, {"title": "B", "snippet": "sb", "url": "https://b.example"}]))
    monkeypatch.setattr(research, "fetch_text", AsyncMock(return_value="long page text " * 30))
    res = await registry.get("research").run(ToolContext(), question="best laptop for students")
    assert res.success and "[1] A" in res.content_for_llm and "[2] B" in res.content_for_llm


def test_find_project(tmp_path, monkeypatch):
    from tools.builtin import dev

    proj = tmp_path / "jarvis-brain-os"
    (proj / ".git").mkdir(parents=True)
    monkeypatch.setattr(dev, "SEARCH_ROOTS", [str(tmp_path)])
    assert dev.find_project("jarvis brain os") == proj
    assert dev.find_project("jarvis") == proj  # fuzzy match on git repos, then remembered


# ── API ──────────────────────────────────────────────────────────────────────
def test_productivity_api(sync_client):
    r = sync_client.post("/reminders", json={"text": "call dad", "when": "in 2 hours"})
    assert r.status_code == 200 and "today" in r.json()["description"] or "tomorrow" in r.json()["description"]
    assert any(x["text"] == "call dad" for x in sync_client.get("/reminders").json()["reminders"])
    assert sync_client.post("/reminders", json={"text": "x", "when": "whenever"}).status_code == 422
    assert sync_client.post("/todos", json={"text": "pay rent"}).status_code == 200
    assert any(t["text"] == "pay rent" for t in sync_client.get("/todos").json()["todos"])
    names = [r["name"] for r in sync_client.get("/routines").json()["routines"]]
    assert "study mode" in names
    page = sync_client.get("/console")
    assert page.status_code == 200 and "Content-Security-Policy" in page.headers and "JARVIS" in page.text


def test_env_example_lists_every_secret():
    text = open(os.path.join(os.path.dirname(__file__), "..", ".env.example"), encoding="utf-8").read()
    for key in ("GEMINI_API_KEY", "GROQ_API_KEY", "OPENAI_API_KEY", "ANTHROPIC_API_KEY", "PRIVACY_MODE", "HOST", "DRY_RUN"):
        assert key in text


@pytest.mark.parametrize("utterance,text", [
    ("remind me in 10 minutes to stretch", "stretch"),
    ("remind me to stretch in 10 minutes", "stretch"),
    ("remind me at 6 pm to submit the report", "submit the report"),
    ("remind me every weekday at 9 to open class links", "open class links"),
    ("set a reminder for tomorrow 9am to call mom", "call mom"),
    ("remind me on friday to pay rent", "pay rent"),
])
def test_reminder_phrasings_take_the_fast_path(utterance, text):
    """Time words after "remind me" ("in", "at", "on", "every") must reach the time parser."""
    from brain.router import fast_route
    from services.timeparse import parse_when, strip_time_phrase

    decision = fast_route(utterance)
    assert decision is not None and decision.tool == "set_reminder"
    when = parse_when(decision.args["when"])
    assert when is not None and strip_time_phrase(decision.args["text"], when) == text
