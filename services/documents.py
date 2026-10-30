"""Document text extraction + local retrieval (BM25) for Q&A over a folder — no cloud needed.

Formats: txt/md/code/csv/json, PDF (pypdf or PyMuPDF), DOCX (python-docx), PPTX (python-pptx),
XLSX (openpyxl). The folder index is cached in data/doc_index/ and refreshed when files change.
"""

from __future__ import annotations

import csv
import hashlib
import json
import math
import os
import re
from collections import Counter
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

from core.paths import data_dir

TEXT_EXTS = {".txt", ".md", ".py", ".js", ".ts", ".json", ".html", ".css", ".java", ".c", ".cpp", ".yaml", ".yml", ".log", ".ini", ".rst"}
DOC_EXTS = TEXT_EXTS | {".pdf", ".docx", ".pptx", ".xlsx", ".csv"}
_WORD = re.compile(r"[a-z0-9]{2,}")
_STOP = set("the a an and or of to in on for is are was were be with by at as it this that from not but you your i".split())


def extract_text(path: Path, max_chars: int = 200_000) -> List[Tuple[str, str]]:
    """Return [(location_label, text)] — e.g. ('p.3', '...') for PDFs, ('slide 2', ...) for PPTX."""
    ext = path.suffix.lower()
    out: List[Tuple[str, str]] = []
    if ext in TEXT_EXTS:
        out.append(("", path.read_text(encoding="utf-8", errors="ignore")[:max_chars]))
    elif ext == ".csv":
        with open(path, newline="", encoding="utf-8", errors="ignore") as fh:
            rows = [", ".join(r) for _, r in zip(range(2000), csv.reader(fh))]
        out.append(("", "\n".join(rows)))
    elif ext == ".pdf":
        # pypdf (BSD) is the declared dependency; PyMuPDF (AGPL) is used only if the user installed it.
        try:
            from pypdf import PdfReader  # type: ignore

            for i, page in enumerate(PdfReader(str(path)).pages, 1):
                out.append((f"p.{i}", page.extract_text() or ""))
        except ImportError:
            import pymupdf  # type: ignore

            with pymupdf.open(str(path)) as doc:
                for i, page in enumerate(doc, 1):
                    out.append((f"p.{i}", page.get_text() or ""))
    elif ext == ".docx":
        import docx  # type: ignore

        d = docx.Document(str(path))
        out.append(("", "\n".join(p.text for p in d.paragraphs)))
        for ti, table in enumerate(d.tables, 1):
            out.append((f"table {ti}", "\n".join(" | ".join(c.text for c in row.cells) for row in table.rows)))
    elif ext == ".pptx":
        from pptx import Presentation  # type: ignore

        for i, slide in enumerate(Presentation(str(path)).slides, 1):
            texts = [sh.text_frame.text for sh in slide.shapes if getattr(sh, "has_text_frame", False)]
            out.append((f"slide {i}", "\n".join(texts)))
    elif ext == ".xlsx":
        from openpyxl import load_workbook  # type: ignore

        wb = load_workbook(str(path), read_only=True, data_only=True)
        for ws in wb.worksheets:
            rows = [" | ".join("" if v is None else str(v) for v in row) for _, row in zip(range(1000), ws.iter_rows(values_only=True))]
            out.append((f"sheet {ws.title}", "\n".join(rows)))
    else:
        raise ValueError(f"Unsupported document type: {ext}")
    total, clipped = 0, []
    for label, text in out:
        if total >= max_chars:
            break
        clipped.append((label, text[: max_chars - total]))
        total += len(text)
    return clipped


def chunk(text: str, size: int = 900, overlap: int = 150) -> List[str]:
    text = re.sub(r"[ \t]+", " ", text).strip()
    if len(text) <= size:
        return [text] if text else []
    chunks, start = [], 0
    while start < len(text):
        end = min(len(text), start + size)
        cut = text.rfind("\n", start + size // 2, end)
        end = cut if cut > start else end
        chunks.append(text[start:end].strip())
        start = max(end - overlap, start + 1)
    return [c for c in chunks if c]


def tokens(text: str) -> List[str]:
    return [w for w in _WORD.findall(text.lower()) if w not in _STOP]


class FolderIndex:
    """BM25 index over a folder's documents, cached on disk keyed by folder path."""

    def __init__(self, folder: Path) -> None:
        self.folder = folder
        key = hashlib.sha1(str(folder.resolve()).encode()).hexdigest()[:16]
        self.cache = data_dir() / "doc_index" / f"{key}.json"
        self.chunks: List[Dict[str, Any]] = []
        self.df: Counter = Counter()
        self.avgdl = 1.0

    def _files(self) -> Dict[str, float]:
        files = {}
        for dirpath, dirnames, filenames in os.walk(self.folder):
            dirnames[:] = [d for d in dirnames if not d.startswith(".") and d not in ("node_modules", "__pycache__")]
            for f in filenames:
                p = Path(dirpath) / f
                if p.suffix.lower() in DOC_EXTS:
                    try:
                        files[str(p)] = p.stat().st_mtime
                    except OSError:
                        pass
            if len(files) > 2000:
                break
        return files

    def build(self) -> int:
        files = self._files()
        cached: Dict[str, Any] = {}
        try:
            cached = json.loads(self.cache.read_text(encoding="utf-8"))
        except (OSError, ValueError):
            pass
        old_files = cached.get("files", {})
        keep = [c for c in cached.get("chunks", []) if old_files.get(c["file"]) == files.get(c["file"])]
        stale = {f for f, m in files.items() if old_files.get(f) != m}
        new_chunks = []
        for f in stale:
            try:
                for label, text in extract_text(Path(f)):
                    for piece in chunk(text):
                        new_chunks.append({"file": f, "loc": label, "text": piece})
            except Exception:
                continue
        self.chunks = keep + new_chunks
        self.cache.parent.mkdir(parents=True, exist_ok=True)
        self.cache.write_text(json.dumps({"files": files, "chunks": self.chunks}), encoding="utf-8")
        self._prepare()
        return len(self.chunks)

    def _prepare(self) -> None:
        self.df = Counter()
        total = 0
        for c in self.chunks:
            toks = tokens(c["text"])
            c["_tf"] = Counter(toks)
            c["_len"] = len(toks)
            total += len(toks)
            self.df.update(set(toks))
        self.avgdl = total / max(len(self.chunks), 1)

    def search(self, query: str, k: int = 6, k1: float = 1.5, b: float = 0.75) -> List[Dict[str, Any]]:
        q = tokens(query)
        n = len(self.chunks)
        scored = []
        for c in self.chunks:
            score = 0.0
            for term in q:
                tf = c["_tf"].get(term, 0)
                if not tf:
                    continue
                idf = math.log(1 + (n - self.df[term] + 0.5) / (self.df[term] + 0.5))
                score += idf * tf * (k1 + 1) / (tf + k1 * (1 - b + b * c["_len"] / self.avgdl))
            if score > 0:
                scored.append((score, c))
        scored.sort(key=lambda x: -x[0])
        return [{"file": c["file"], "loc": c["loc"], "text": c["text"], "score": round(s, 3)} for s, c in scored[:k]]


def ask_folder(folder: Path, question: str, k: int = 6) -> Tuple[List[Dict[str, Any]], int]:
    index = FolderIndex(folder)
    count = index.build()
    return index.search(question, k), count


def extract_tables(path: Path) -> List[List[List[str]]]:
    ext = path.suffix.lower()
    tables: List[List[List[str]]] = []
    if ext == ".docx":
        import docx  # type: ignore

        for t in docx.Document(str(path)).tables:
            tables.append([[c.text.strip() for c in row.cells] for row in t.rows])
    elif ext == ".pdf":
        try:
            import pdfplumber  # type: ignore
        except ImportError:
            raise RuntimeError("PDF table extraction needs 'pdfplumber' (pip install pdfplumber).")
        with pdfplumber.open(str(path)) as pdf:
            for page in pdf.pages:
                for t in page.extract_tables():
                    tables.append([[("" if c is None else str(c)).strip() for c in row] for row in t])
    elif ext in (".md", ".txt"):
        current: List[List[str]] = []
        for line in path.read_text(encoding="utf-8", errors="ignore").splitlines():
            if line.strip().startswith("|") and line.count("|") >= 2:
                cells = [c.strip() for c in line.strip().strip("|").split("|")]
                if not all(set(c) <= set("-: ") for c in cells):
                    current.append(cells)
            elif current:
                tables.append(current)
                current = []
        if current:
            tables.append(current)
    elif ext == ".csv":
        with open(path, newline="", encoding="utf-8", errors="ignore") as fh:
            tables.append([row for row in csv.reader(fh)])
    else:
        raise RuntimeError(f"I can't extract tables from {ext} files.")
    return tables


def write_xlsx(tables: List[List[List[str]]], dest: Path) -> None:
    from openpyxl import Workbook  # type: ignore

    wb = Workbook()
    if wb.active is not None:
        wb.remove(wb.active)  # type: ignore[arg-type]
    for i, table in enumerate(tables, 1):
        ws = wb.create_sheet(f"Table {i}")
        for row in table:
            ws.append(row)
    wb.save(str(dest))


def convert(src: Path, target_ext: str, dest: Optional[Path] = None) -> Path:
    target_ext = "." + target_ext.lower().lstrip(".")
    dest = dest or src.with_suffix(target_ext)
    sext = src.suffix.lower()
    if target_ext in (".txt", ".md") and sext in DOC_EXTS:
        parts = extract_text(src)
        body = "\n\n".join((f"## {label}\n{text}" if label and target_ext == ".md" else text) for label, text in parts)
        dest.write_text(body, encoding="utf-8")
    elif target_ext == ".docx" and sext in (".txt", ".md"):
        import docx  # type: ignore

        d = docx.Document()
        for line in src.read_text(encoding="utf-8", errors="ignore").splitlines():
            m = re.match(r"^(#{1,3})\s+(.*)", line)
            if m:
                d.add_heading(m.group(2), level=len(m.group(1)))
            elif line.strip():
                d.add_paragraph(line)
        d.save(str(dest))
    elif target_ext == ".xlsx" and sext == ".csv":
        write_xlsx(extract_tables(src), dest)
    elif target_ext == ".csv" and sext == ".xlsx":
        from openpyxl import load_workbook  # type: ignore

        ws = load_workbook(str(src), read_only=True, data_only=True).worksheets[0]
        with open(dest, "w", newline="", encoding="utf-8") as fh:
            writer = csv.writer(fh)
            for row in ws.iter_rows(values_only=True):
                writer.writerow(["" if v is None else v for v in row])
    else:
        raise RuntimeError(f"Converting {sext} → {target_ext} isn't supported yet.")
    return dest
